"""Documents router: extraction helpers (real fixtures) + upload/preview/raw endpoints."""
import sys
from pathlib import Path

import allure
import pytest

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

from ui.backend.routers import documents as d  # noqa: E402


def _make_docx(path: Path, text: str):
    from docx import Document
    doc = Document()
    doc.add_paragraph(text)
    doc.save(str(path))


def _make_xlsx(path: Path, rows):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    wb.save(str(path))


@allure.feature("Pipeline API")
@allure.story("Documents extraction")
class TestExtraction:
    @allure.title("_extract_docx_text reads paragraph text from a real .docx")
    def test_docx_text(self, tmp_path):
        f = tmp_path / "a.docx"; _make_docx(f, "Hello compliance world")
        content, truncated = d._extract_docx_text(f, 10000)
        assert "Hello compliance world" in content
        assert truncated is False

    @allure.title("_extract_xlsx_preview renders sheet cells from a real .xlsx")
    def test_xlsx_preview(self, tmp_path):
        f = tmp_path / "b.xlsx"; _make_xlsx(f, [["Rule", "Risk"], ["LTV<=80", "high"]])
        content, _ = d._extract_xlsx_preview(f, 10000)
        assert "Rule" in content and "LTV<=80" in content

    @allure.title("_extract_pdf_text returns '' for a non-PDF without raising")
    def test_pdf_text_graceful(self, tmp_path):
        f = tmp_path / "x.pdf"; f.write_bytes(b"not really a pdf")
        assert d._extract_pdf_text(f) == ""


@allure.feature("Pipeline API")
@allure.story("Documents upload / preview / raw")
class TestUploadPreview:
    @pytest.fixture
    def client(self, tmp_path, monkeypatch):
        monkeypatch.setattr(d, "PROJECT_ROOT", tmp_path)
        (tmp_path / "compliance-files").mkdir()
        from fastapi.testclient import TestClient
        from ui.backend.main import app
        return TestClient(app), tmp_path

    @allure.title("Upload a text doc into a subdir, then preview + raw it")
    def test_upload_then_preview(self, client):
        c, root = client
        c.post("/api/documents/folder", json={"name": "mortgage"})
        up = c.post(
            "/api/documents/upload?subdir=mortgage",
            files=[("files", ("rule.txt", b"The maximum LTV is 80 percent.", "text/plain"))],
        )
        assert up.status_code == 200
        assert (root / "compliance-files" / "mortgage" / "rule.txt").exists()

        prev = c.get("/api/documents/preview/mortgage/rule.txt")
        assert prev.status_code == 200
        assert "maximum LTV" in prev.json()["content"]

        raw = c.get("/api/documents/raw/mortgage/rule.txt")
        assert raw.status_code == 200

    @allure.title("Preview of a .docx returns docx content")
    def test_preview_docx(self, client):
        c, root = client
        (root / "compliance-files" / "d1").mkdir()
        _make_docx(root / "compliance-files" / "d1" / "policy.docx", "Borrower must verify income")
        prev = c.get("/api/documents/preview/d1/policy.docx")
        assert prev.status_code == 200
        assert "Borrower must verify income" in prev.json()["content"]

    @allure.title("Preview traversal attempt is rejected")
    def test_preview_traversal(self, client):
        c, _ = client
        assert c.get("/api/documents/preview/x/..%2f..%2fconfig.json").status_code in (400, 404)
